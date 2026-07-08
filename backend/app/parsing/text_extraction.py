"""
Extracts raw text from an uploaded resume file.

Supports .pdf, .docx and .txt. PDF/DOCX extraction depends on `pypdf` and
`python-docx` respectively (see requirements.txt) - both are free, pure
open-source libraries with no model downloads or API keys involved.

Deliberately NOT supported: scanned/image-only PDFs (no OCR). If a PDF
extracts to near-empty text, we raise ExtractionError with a message that
explains why, instead of silently scoring a blank resume.
"""

from __future__ import annotations
import io


class ExtractionError(Exception):
    """Raised when a file can't be turned into usable text."""


def extract_text(filename: str, content: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "txt":
        return _extract_txt(content)
    if ext == "pdf":
        return _extract_pdf(content)
    if ext == "docx":
        return _extract_docx(content)

    raise ExtractionError(
        f"Unsupported file type '.{ext}'. Please upload a .pdf, .docx, or .txt resume."
    )


def _extract_txt(content: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode this .txt file as text.")


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError(
            "PDF support requires the 'pypdf' package. Run: pip install pypdf"
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(content))
        pages_text = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # pypdf raises several distinct exception types
        raise ExtractionError(f"Could not read this PDF ({exc}).") from exc

    text = "\n".join(pages_text).strip()
    if len(text) < 30:
        raise ExtractionError(
            "This PDF appears to have little to no extractable text. "
            "It may be a scanned image rather than a text-based PDF "
            "(OCR is not supported) - try exporting it as text-based PDF, "
            "or upload a .docx / .txt version instead."
        )
    return text


def _extract_docx(content: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise ExtractionError(
            "DOCX support requires the 'python-docx' package. Run: pip install python-docx"
        ) from exc

    try:
        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    except Exception as exc:
        raise ExtractionError(f"Could not read this DOCX file ({exc}).") from exc

    text = "\n".join(p for p in parts if p.strip())
    if len(text) < 30:
        raise ExtractionError("This DOCX file appears to contain little to no text.")
    return text
